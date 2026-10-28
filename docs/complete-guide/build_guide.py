from pathlib import Path
import json,sys,re
from datetime import datetime
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from guide_assets import create_assets
from guide_content import PAGES,p,fig,table
from guide_appendices import appendices,GLOSSARY

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
manifest=json.loads((ROOT/'reports/source-manifest.json').read_text('utf-8-sig'))
checks=json.loads((ROOT/'reports/verification-results.json').read_text('utf-8-sig'))
assert len(checks)==67 and all(c['Passed'] for c in checks)
counts=json.loads(checks[-1]['Detail'])
assets=create_assets()
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.66)
sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
sec.different_first_page_header_footer=True
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Header','Footer']:
 st=styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.font.underline=False
for name,size in [('Normal',11),('Title',27),('Subtitle',16),('Heading 1',19),('Heading 2',12.5),('Heading 3',11.5),('Caption',9),('Header',8),('Footer',9)]:styles[name].font.size=Pt(size)
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.1
styles['Normal'].paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2','Heading 3']:
 styles[name].paragraph_format.keep_with_next=True
 styles[name].paragraph_format.space_before=Pt(9)
 styles[name].paragraph_format.space_after=Pt(7)
styles['Heading 1'].paragraph_format.space_before=Pt(0)
styles['Caption'].paragraph_format.space_before=Pt(3)
styles['Caption'].paragraph_format.space_after=Pt(8)
styles['Caption'].font.italic=False
for name in ['Title','Subtitle']:
 pr=styles[name].element.get_or_add_pPr()
 for x in list(pr.findall(qn('w:pBdr'))):pr.remove(x)
 borders=OxmlElement('w:pBdr')
 for side in ['top','bottom','left','right','between']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');borders.append(e)
 pr.append(borders)
header=sec.header.paragraphs[0];header.text='HealthRetailers complete implementation guide'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='HealthRetailers complete project guide'
doc.core_properties.subject='Dirty source exports to a tested SQL Server warehouse using native SSIS packages'
doc.core_properties.author='HealthRetailers project group'
doc.core_properties.last_modified_by='HealthRetailers project group'
doc.core_properties.keywords='HealthRetailers SQL Server SSIS staging data warehouse stock purchases'

def clean(s):return str(s).replace('\u2011','-').replace('\u2013','-').replace('\u2014','-')
def para(s,style=None):
 p=doc.add_paragraph(clean(s),style)
 p.paragraph_format.widow_control=True
 return p
def format_table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[6.85/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 pr.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for side,val in [('top',65),('bottom',65),('left',85),('right',85)]:
  e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
 pr.append(margins)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=clean(h)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=clean(v)
 for ri,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'233D56' if ri==0 else ('F2F5F8' if ri%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shd)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1.05
    if ci==len(headers)-1 and (len(str(cell.text))<18):pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in pp.runs:
     run.font.name='Arial';run.font.size=Pt(9.4);run.bold=ri==0
     run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
 para('').paragraph_format.space_after=Pt(1)
 return t

figure_count=0
def add_figure(name,caption,width=6.85,maxheight=4.8):
 global figure_count
 file=assets/(name+'.png');im=Image.open(file);w,h=im.size
 width=min(width,maxheight*w/h)
 pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.keep_with_next=True
 pp.add_run().add_picture(str(file),width=Inches(width))
 figure_count+=1
 para(f'Figure {figure_count}. '+caption,'Caption')

def block(b):
 if b[0]=='p':para(b[1])
 elif b[0]=='h':para(b[1],'Heading 2')
 elif b[0]=='fig':add_figure(*b[1:])
 elif b[0]=='table':format_table(*b[1:])
 elif b[0]=='code':
  pp=doc.add_paragraph();pp.paragraph_format.left_indent=Inches(.12)
  pp.paragraph_format.right_indent=Inches(.05);pp.paragraph_format.line_spacing=1.05
  pp.paragraph_format.space_after=Pt(9)
  pp.paragraph_format.keep_together=True
  rr=pp.add_run(clean(b[1]));rr.font.name='Consolas';rr.font.size=Pt(9)
 elif b[0]=='sources':
  rows=[(s['source'],s['table'],f"{s['rows']:,}",s['encoding'].replace('utf-8-sig','UTF8').upper()) for s in manifest]
  rows.append(('TOTAL','17 raw staging tables',f"{sum(s['rows'] for s in manifest):,}",''))
  format_table(['Source export','Staging table','Rows','Decoding'],rows,[2.4,2.5,.75,1.2])
 elif b[0]=='results':
  order=['dim_produit','dim_client','dim_fournisseur','dim_paiement','dim_temps','fait_stock_actuel','fait_achats']
  format_table(['Warehouse table','Loaded rows'],[(n,f'{counts[n]:,}') for n in order],[4.8,2.05])
 else:raise ValueError(b[0])

# Cover with a useful scope statement, rather than a decorative blank page.
para('HealthRetailers data warehouse','Title')
para('Complete implementation guide','Subtitle')
para('English edition  |  Final implemented scope  |  28 October 2026')
para('From dirty source exports to a tested SQL Server warehouse')
para('This guide explains the complete implemented project so that a group member can understand, demonstrate and rerun it. It follows the actual data through file preparation, staging, dimensions, facts and validation. Technical terms are defined as they appear and collected in a glossary.')
add_figure('journey','The implemented pipeline at a glance.',6.85,1.5)
para('The delivered scope is five dimensions and two facts: current stock and validated supplier invoice lines. All 67 recorded checks passed. The guide distinguishes raw evidence, derived information, known unknowns and out-of-scope business processes.')
format_table(['What is ready','Verified result'],[
 ('Raw source integration','17 selected exports and 244300 staged rows'),('Dimension foundation','Product  client  supplier  payment mode  calendar'),('Fact loading','6842 stock products and 6964 purchase lines'),('Purchase coverage','880 validated invoices with exact net-header reconciliation'),('Technical evidence','Native package runs  stable reruns  tested repairs  trusted FKs'),
 ],[2.3,4.55])
para('Reading copies: this PDF and the editable Word document contain the same guide. The operational package and SQL files remain the execution authority.')
doc.add_page_break()
para('Contents','Title')
pp=doc.add_paragraph()
begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\o "1-1" \\h \\z \\u '
sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
txt=OxmlElement('w:t');txt.text='Contents is updated during final rendering.'
end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
for el in [begin,instr,sep,txt,end]:
 rr=OxmlElement('w:r');rr.append(el);pp._p.append(rr)
for title,blocks in PAGES+appendices(manifest,checks):
 doc.add_page_break();para(title,'Heading 1')
 for b in blocks:block(b)

settings=doc.settings.element
uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');settings.append(uf)
out=BASE/'HealthRetailers_Complete_Project_Guide_EN.docx';doc.save(out)
text='\n'.join(p.text for p in doc.paragraphs)
meta={'docx':str(out),'chapters':len(PAGES),'glossaryTerms':len(GLOSSARY),'figures':figure_count,'sourceTables':len(manifest),'checks':len(checks),'narrativeWordCount':len(text.split()),'buildTime':datetime.now().isoformat(timespec='seconds')}
(BASE/'guide-build.json').write_text(json.dumps(meta,indent=2),encoding='utf-8')
print(json.dumps(meta,indent=2))
