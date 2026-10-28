from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
import hashlib,json
from lxml import etree as E
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from guide_assets import create_assets
from guide_appendices import TECHNICAL_STAGING

base=Path(__file__).resolve().parent
source=base/'HealthRetailers_Complete_Project_Guide_EN.docx'
old={hashlib.sha256(p.read_bytes()).hexdigest():p for p in (base/'assets').glob('*.png')}
create_assets()
replace={key:p.read_bytes() for key,p in old.items() if hashlib.sha256(p.read_bytes()).hexdigest()!=key}
with ZipFile(source) as z:data={n:z.read(n) for n in z.namelist()}
for n,b in list(data.items()):
 if n.startswith('word/media/') and hashlib.sha256(b).hexdigest() in replace:data[n]=replace[hashlib.sha256(b).hexdigest()]
with ZipFile(source,'w',8) as z:
 for n,b in data.items():z.writestr(n,b)

d=Document(source)
def no_border(pr):
 for x in list(pr.findall(qn('w:pBdr'))):pr.remove(x)
 borders=OxmlElement('w:pBdr')
 for side in ['top','left','bottom','right','between']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');borders.append(e)
 pr.append(borders)
for name in ['Title','Subtitle']:no_border(d.styles[name].element.get_or_add_pPr())
for pp in d.paragraphs:
 if pp.style.name in ['Title','Subtitle']:no_border(pp._p.get_or_add_pPr())

anchor=next(pp._p for pp in d.paragraphs if pp.text=='Appendix C Recorded checks 1')
# Keep the appendix page break after the new technical description, not before it.
before=anchor.getprevious()
if before is not None and before.xpath('.//w:br'):anchor=before
for block in TECHNICAL_STAGING:
 if block[0] in ['h','p']:
  pp=d.add_paragraph(block[1],'Heading 2' if block[0]=='h' else 'Normal')
  anchor.addprevious(pp._p)
 elif block[0]=='table':
  headers,rows,widths=block[1:];t=d.add_table(rows=1,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
  for i,x in enumerate(headers):t.rows[0].cells[i].text=x
  for row in rows:
   cells=t.add_row().cells
   for c,v in zip(cells,row):c.text=v
  repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
  pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
  for side in ['top','left','bottom','right','insideH','insideV']:
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  pr.append(borders);mar=OxmlElement('w:tblCellMar')
  for side in ['top','bottom','left','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
  pr.append(mar)
  for ri,row in enumerate(t.rows):
   el=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(el)
   for ci,cell in enumerate(row.cells):
    cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'233D56' if ri==0 else ('F2F5F8' if ri%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(sh)
    for pp in cell.paragraphs:
     pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.line_spacing=1.05
     for rr in pp.runs:rr.font.name='Arial';rr.font.size=Pt(9.4);rr.bold=ri==0;rr.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
  anchor.addprevious(t._tbl)
d.save(source)

# Produce native Word rendering copies while keeping the full editable source intact.
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};W='{'+ns['w']+'}'
with ZipFile(source) as z:data={n:z.read(n) for n in z.namelist()}
r=E.fromstring(data['word/document.xml']);body=r.find(W+'body');sect=deepcopy(body.find(W+'sectPr'));groups=[];group=[]
for child in body:
 if child.tag==W+'sectPr':continue
 if child.xpath('.//w:br[@w:type="page"]',namespaces=ns):groups.append(group);group=[]
 else:group.append(deepcopy(child))
if group:groups.append(group)
items=[]
for i,group in enumerate(groups,1):
 tree=deepcopy(r);b=tree.find(W+'body');b.clear()
 for el in group:b.append(el)
 s=deepcopy(sect)
 if i>1:
  for x in s.findall(W+'titlePg'):s.remove(x)
 b.append(s)
 for el in list(tree.xpath('//w:fldChar|//w:instrText',namespaces=ns)):el.getparent().remove(el)
 settings=E.fromstring(data['word/settings.xml'])
 for x in settings.findall(W+'updateFields'):settings.remove(x)
 file=base/'qa/word-parts'/f'part-{i:02}.docx'
 with ZipFile(file,'w',8) as out:
  for name,value in data.items():
   if name=='word/document.xml':value=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
   if name=='word/settings.xml':value=E.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
   out.writestr(name,value)
 title=''.join(b.xpath('.//w:p',namespaces=ns)[0].xpath('.//w:t/text()',namespaces=ns))
 items.append({'part':i,'docx':str(file),'title':title})
(base/'qa/parts.json').write_text(json.dumps(items,indent=2),encoding='utf-8')
print('Final layout revision complete; full field inventory retained.')
